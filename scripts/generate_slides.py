import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

OUTPUT_DIR = Path(
"data/presentations"
)

def parse_sections(markdown):
sections = []

current_title = None
current_lines = []

for line in markdown.splitlines():

    if line.startswith("## "):

        if current_title:
            sections.append(
                (
                    current_title,
                    current_lines
                )
            )

        current_title = (
            line[3:].strip()
        )

        current_lines = []

    elif current_title:
        current_lines.append(
            line
        )

if current_title:
    sections.append(
        (
            current_title,
            current_lines
        )
    )

return sections


def add_title_slide(prs, title, subtitle):
slide = prs.slides.add_slide(
prs.slide_layouts[0]
)

slide.shapes.title.text = title

slide.placeholders[1].text = (
    subtitle
)

return slide


def add_content_slide(prs, title, lines):
slide = prs.slides.add_slide(
prs.slide_layouts[1]
)

slide.shapes.title.text = title

text_frame = (
    slide.placeholders[1]
    .text_frame
)

text_frame.clear()

first = True

for line in lines:

    line = line.strip()

    if not line:
        continue

    if line.startswith("- "):
        text = line[2:].strip()
    elif line.startswith("* "):
        text = line[2:].strip()
    elif line.startswith("1. "):
        text = line[3:].strip()
    else:
        text = line

    if first:
        paragraph = text_frame.paragraphs[0]
        first = False
    else:
        paragraph = text_frame.add_paragraph()

    paragraph.text = text
    paragraph.font.size = Pt(20)

return slide


def main():
if len(sys.argv) < 2:
print(
"Usage:"
)
print(
"python scripts/generate_slides.py "
"<summary.md>"
)
sys.exit(1)

summary_path = Path(
    sys.argv[1]
)

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True
)

markdown = summary_path.read_text(
    encoding="utf-8"
)

sections = parse_sections(
    markdown
)

prs = Presentation()

prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

date = summary_path.stem.replace(
    "-summary",
    ""
)

add_title_slide(
    prs,
    "Daily Knowledge Report",
    date
)

for title, lines in sections:

    if title.lower() == "sources":
        continue

    add_content_slide(
        prs,
        title,
        lines
    )

output = (
    OUTPUT_DIR
    / f"{date}.pptx"
)

prs.save(output)

print(
    f"Generated presentation: {output}"
)


if name == "main":
main()
